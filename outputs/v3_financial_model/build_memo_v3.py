from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT="SVT_Nexus_2026_2029_Financial_Estimate_Memorandum_V3.docx"
NAVY="16324F"; TEAL="1B7F79"; PALE="EAF4F3"; ORANGE="FFF4E5"; GRAY="667085"

def shade(c,color):
    p=c._tc.get_or_add_tcPr(); x=OxmlElement('w:shd'); x.set(qn('w:fill'),color); p.append(x)
def pad(c,v=110):
    p=c._tc.get_or_add_tcPr(); m=OxmlElement('w:tcMar')
    for k in ('top','start','bottom','end'):
        x=OxmlElement('w:'+k); x.set(qn('w:w'),str(v)); x.set(qn('w:type'),'dxa'); m.append(x)
    p.append(m)
def repeat(row):
    p=row._tr.get_or_add_trPr(); x=OxmlElement('w:tblHeader'); x.set(qn('w:val'),'true'); p.append(x)
def table(doc,heads,rows,widths=None):
    t=doc.add_table(rows=1,cols=len(heads)); t.style='Table Grid'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
    for i,h in enumerate(heads):
        c=t.rows[0].cells[i]; c.text=str(h); shade(c,TEAL); pad(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(8.2)
    repeat(t.rows[0])
    for n,row in enumerate(rows):
        cells=t.add_row().cells
        for i,v in enumerate(row):
            cells[i].text=str(v); pad(cells[i]); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if n%2: shade(cells[i],'F4F7FA')
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after=Pt(0)
                for r in p.runs: r.font.name='Aptos'; r.font.size=Pt(8.2); r.font.color.rgb=RGBColor(35,45,55)
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t
def callout(doc,text,color=ORANGE):
    t=doc.add_table(rows=1,cols=1); c=t.cell(0,0); shade(c,color); pad(c,170)
    p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0); r=p.add_run(text); r.bold=True; r.font.color.rgb=RGBColor.from_string(NAVY); r.font.size=Pt(9)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
def bullet(doc,text): doc.add_paragraph(text,style='List Bullet')

d=Document(); sec=d.sections[0]; sec.top_margin=Inches(.68); sec.bottom_margin=Inches(.62); sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
normal=d.styles['Normal']; normal.font.name='Aptos'; normal.font.size=Pt(9.4); normal.font.color.rgb=RGBColor(35,45,55); normal.paragraph_format.space_after=Pt(5); normal.paragraph_format.line_spacing=1.08
for name,size,color in [('Title',28,NAVY),('Heading 1',17,NAVY),('Heading 2',12,TEAL),('Heading 3',10,NAVY)]:
    s=d.styles[name]; s.font.name='Aptos Display'; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color); s.paragraph_format.space_before=Pt(10); s.paragraph_format.space_after=Pt(5); s.paragraph_format.keep_with_next=True
hp=sec.header.paragraphs[0]; hp.text="SVT NEXUS  |  VERSION 3 FINANCIAL SUPPORT MEMORANDUM"; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
for r in hp.runs: r.font.size=Pt(7.2); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(TEAL)
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=fp.add_run("CONFIDENTIAL DRAFT — COUNSEL AND ADVISER REVIEW REQUIRED  |  21 AUGUST 2026  |  "); r.font.size=Pt(7); r.font.color.rgb=RGBColor.from_string(GRAY); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); fp._p.append(fld)

p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(85); r=p.add_run("SVT NEXUS"); r.font.name='Aptos Display'; r.font.size=Pt(34); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(NAVY)
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("2026–2029 Financial Estimate and\nOffering Support Memorandum"); r.font.name='Aptos Display'; r.font.size=Pt(22); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(TEAL)
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(14); r=p.add_run("VERSION 3  •  PROPOSED HYBRID CONVERTIBLE NOTE CASE"); r.font.size=Pt(11); r.font.bold=True; r.font.color.rgb=RGBColor.from_string(GRAY)
callout(d,"$3.0M maximum • $1.25M Phase 1 / $1.75M milestone-gated Phase 2 • 4% quarterly cash coupon • 5% PIK • illustrative conversion terms",PALE)
p=d.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(75); r=p.add_run("This document is internal financial support, not a PPM, offering circular, legal opinion, or recommendation to invest."); r.font.size=Pt(8.5); r.font.italic=True; r.font.color.rgb=RGBColor.from_string(GRAY)
d.add_page_break()

d.add_heading('1. Executive Decision Summary',1)
d.add_paragraph("Version 3 responds to the capital-efficiency, debt-service, growth-capacity, sales-evidence, tax, accounting, and offering-compliance concerns raised in the Version 2 review. It replaces the all-cash 9% coupon with a proposed 4% cash / 5% PIK structure, separates the raise into two phases, expands the operating investment supporting enterprise acquisition, and connects the Base Case to 12, 18, and 24 ending clients in 2027, 2028, and 2029.")
callout(d,"Proposed terms are modeling assumptions only. Management and securities counsel must approve the actual note, conversion mechanics, offering exemption, PPM, subscription agreement, and closing process.")
table(d,["Metric","2026","2027","2028","2029"],[
 ["Note proceeds","$150,000","$2,850,000","—","—"],["Gross revenue","—","$1,392,000","$2,765,070","$4,055,436"],["Operating income","$(42,500)","$(451,940)","$440,183","$1,266,734"],["Total interest expense","$(1,147)","$(167,042)","$(270,740)","$(270,000)"],["Net income","$(43,647)","$(618,982)","$89,330","$737,584"],["Ending cash","$106,990","$2,430,809","$2,670,550","$3,558,133"],["Note carrying amount","$150,637","$3,093,438","$3,243,849","$3,393,849"]
],[2.15,1.05,1.05,1.05,1.05])
d.add_paragraph("The higher 2027 operating loss reflects deliberate deployment into sales, personnel, and technology. The 2028–2029 improvement depends on the modeled client expansion and must be supported by pipeline evidence and delivery capacity.")

d.add_heading('2. Proposed Financing Architecture',1)
d.add_heading('2.1 Phased funding',2)
table(d,["Phase","Maximum","Release condition","Modeled schedule"],[
 ["Phase 1","$1,250,000","Initial runway and 12-client build","December 2026–April 2027"],["Phase 2","$1,750,000","Prior-month annualized gross revenue of at least $1.0M","May–November 2027 if achieved"],["Maximum","$3,000,000","Both phases close","Subject to actual subscriptions and receipts"]
],[1.0,1.25,3.2,1.65])
d.add_paragraph("The Base Case reaches a $1.2 million annualized gross-revenue run-rate in April 2027, so the spreadsheet releases the first Phase 2 tranche in May. This is a model gate, not a legal funding commitment; definitive documents must specify whether the second phase is optional, committed, warrant-backed, or separately offered.")
d.add_heading('2.2 Cash and PIK interest',2)
bullet(d,"Cash coupon: 4% per annum, accrued from each actual funding date using an ACT/365 planning convention and paid quarterly.")
bullet(d,"PIK coupon: 5% per annum, accrued on funded face principal and added to the note carrying amount; the planning model does not compound PIK-on-PIK.")
bullet(d,"Total modeled coupon expense remains 9% before issuance-cost, discount, derivative, or effective-interest adjustments.")
bullet(d,"The 2029 ending carrying amount is approximately $3.394 million, including approximately $394,000 of cumulative PIK.")
d.add_heading('2.3 Illustrative conversion terms',2)
table(d,["Term","Illustrative Version 3 assumption","Required action"],[
 ["Conversion discount","20%","Counsel to define eligible financing price and mechanics"],["Valuation cap","$12.0 million pre-money","Confirm capitalization basis and fully diluted definition"],["Qualified financing","$5.0 million new-money threshold","Define included securities and closing aggregation"],["Maturity","November 30, 2030","Proposed 36 months after final scheduled close"],["Forecast conversion","None through December 2029","No qualifying financing is assumed in the Base Case"]
],[1.65,2.2,3.15])
d.add_paragraph("The conversion terms provide an equity-upside narrative without recording a conversion event that is not supported by the forecast. They must not be presented as finalized until reflected consistently in the term sheet, note, PPM, subscription agreement, capitalization table, and accounting analysis.")

d.add_heading('3. Capital Deployment and Operating Capacity',1)
table(d,["Use","Amount","% of maximum"],[
 ["Personnel (non-CEO)","$600,000","20.0%"],["Sales and enterprise marketing","$420,000","14.0%"],["Technology / AI / cloud","$180,000","6.0%"],["CEO compensation","$120,000","4.0%"],["Legal / compliance / audit","$80,000","2.7%"],["Insurance, travel, office and consultants","$140,000","4.7%"],["Liquidity reserve","$300,000","10.0%"],["Working capital / contingency","$1,160,000","38.7%"],["Total","$3,000,000","100.0%"]
],[3.8,1.35,1.2])
d.add_paragraph("Personnel, enterprise sales, and technology together represent $1.2 million, or 40% of the maximum raise. The sales allocation includes the requested $300,000 increase intended to support two enterprise account executives, one sales-development representative, trade shows, paid acquisition, and pilot demonstrations. Actual compensation plans and hiring dates remain to be documented.")
d.add_paragraph("The remaining working-capital amount is still material. Management should convert it into a board-approved deployment schedule with hiring gates, product milestones, customer-success capacity, and downside controls before investor circulation.")

d.add_heading('4. Sales Funnel and Revenue Support',1)
table(d,["Funnel metric","2027","2028","2029"],[
 ["Qualified leads","240","180","180"],["Demos / discovery","72","54","54"],["Paid or documented pilots","24","18","18"],["New clients","12","6","6"],["Ending clients","12","18","24"],["Lead-to-demo","30.0%","30.0%","30.0%"],["Demo-to-pilot","33.3%","33.3%","33.3%"],["Pilot-to-client","50.0%","33.3%","33.3%"],["Sales spend per new client","$35,000","$77,000","$84,700"]
],[2.8,1.2,1.2,1.2])
d.add_paragraph("These funnel counts are analytical assumptions, not existing pipeline evidence. Before investor use, management should reconcile them to CRM exports, dated proposals, pilot budgets, LOIs/MOUs, definitive contracts, responsible sales owners, expected closing dates, and probability-weighted values.")
d.add_heading('4.1 Revenue recognition',2)
d.add_paragraph("Monthly forecast billing does not establish U.S. GAAP revenue recognition. Each retainer, commission, pilot, setup fee, implementation milestone, variable-payment feature, renewal, cancellation right, and performance obligation must be evaluated under ASC 606. The final accounting may recognize revenue over time, at a point in time, or defer amounts relative to forecast billing.")

d.add_heading('5. Profitability, Cash Flow, and Capital Efficiency',1)
d.add_paragraph("Version 3 intentionally accepts a larger 2027 operating loss to build enterprise acquisition and delivery capacity. The Base Case then increases gross revenue to approximately $2.77 million in 2028 and $4.06 million in 2029, producing positive operating income and net income in both years despite the full 9% combined cash/PIK expense.")
bullet(d,"Cash interest is lower than Version 2 because only the 4% component is paid quarterly.")
bullet(d,"PIK preserves operating cash but increases the amount owed or converted at maturity.")
bullet(d,"Operating cash flow excluding financing remains separately visible in the spreadsheet.")
bullet(d,"Ending cash remains substantial; management should justify it through staged commitments, downside runway, acquisition capacity, and board-controlled deployment—not generic contingency language.")
d.add_heading('5.1 Break-even',2)
d.add_paragraph("At the post-July 2027 twelve-client mix, the model generates $147,000 of gross monthly revenue and approximately a 77.2% contribution margin after bad debt and differentiated direct costs. With the expanded 2027 monthly fixed operating budget of approximately $128,300, indicative gross-revenue break-even is approximately $166,300, or about 13.6 equivalent clients. The 2027 Base Case therefore builds toward, but does not maintain, the expanded-cost break-even level; the 2028 client expansion is critical.")

d.add_heading('6. Tax Stress Test — Section 163(j)',1)
d.add_paragraph("Version 3 adds a simplified stress case that limits current deductible business interest to 30% of modeled adjusted taxable income when the toggle is active. The model produces no current tax provision in 2026 or 2027, approximately $80,113 in 2028, and approximately $259,151 in 2029 under that simplified convention.")
callout(d,"This is not a tax conclusion. Certain small businesses may be exempt, and the legal entity, gross-receipts test, ATI adjustments, PIK/OID treatment, carryforwards, state taxes, and estimated-payment rules require tax-adviser review.")
d.add_paragraph("IRS source: https://www.irs.gov/newsroom/questions-and-answers-about-the-limitation-on-the-deduction-for-business-interest-expense")

d.add_heading('7. Offering and Accounting Compliance Roadmap',1)
table(d,["Area","Version 3 disclosure","Required owner/action"],[
 ["Rule 506(c)","Working exemption only; all purchasers must be accredited and the issuer must take reasonable verification steps","Securities counsel to approve exemption and verification process"],["Form D","Generally due within 15 days after first sale; no filing is represented as completed","Counsel to control EDGAR and state notice calendar"],["PPM package","This memorandum is not a PPM and must not be distributed alone","Prepare PPM, subscription agreement, note, legends and risk factors"],["ASC 606","Revenue schedule is planning-only","Accounting adviser to evaluate executed customer contracts"],["Convertible accounting","Cash/PIK debt is modeled; no derivative liability or issuance-cost amortization is recorded","Accounting adviser to evaluate ASC 470/480/815 and effective interest"],["Data room","Funnel assumptions are not evidence","Management to assemble CRM support, LOIs/MOUs, contracts, budgets and approvals"]
],[1.4,3.15,2.6])
d.add_paragraph("SEC sources: https://www.sec.gov/resources-small-businesses/exempt-offerings/general-solicitation-rule-506c ; https://www.sec.gov/resources-small-businesses/capital-raising-building-blocks/assessing-accredited-investors-under-regulation-d ; https://www.sec.gov/resources-small-businesses/capital-raising-building-blocks/what-form-d")
d.add_paragraph("FASB sources: https://storage.fasb.org/Rev_Rec_Implementation_QAs.pdf ; https://storage.fasb.org/ASU_2020-06.pdf")

d.add_heading('8. Required Documents and Governance Before Launch',1)
table(d,["Phase","Action","Responsible party","Priority"],[
 ["Legal","Draft and reconcile the PPM, subscription agreement, promissory note, Rule 506(c) legends, risk factors, and state notices","Securities counsel","Immediate"],["Terms","Approve cash/PIK coupon, phase mechanics, discount, cap, financing trigger, maturity, defaults, security and subordination","Management and counsel","Immediate"],["Financial","Validate formulas, note accounting, issuance costs, derivative analysis, taxes and disclosures","CPA / accounting adviser","High"],["Sales","Reconcile the funnel to CRM records and define owners, dates, probabilities and evidence","Chief Commercial Officer","High"],["Operations","Approve hiring, technology, delivery-capacity and customer-success gates for 12/18/24 clients","Management / board","High"],["Diligence","Create a controlled data room containing corporate, financial, customer, IP, personnel and offering records","Management and counsel","High"]
],[.85,3.35,1.55,.75])

d.add_heading('9. Forward-Looking Statement and Limitation',1)
d.add_paragraph("The projections are forward-looking management estimates, not guarantees. They depend on successful capital raising, milestone satisfaction, enterprise lead generation, conversion rates, contract execution, pricing, client retention, collectibility, delivery capacity, cost control, tax treatment, financing terms, and continued operations. SVT Nexus is pre-operating for the proposed activity and does not have historical results demonstrating achievement of the projections. Actual results may differ materially.")
d.add_paragraph("This memorandum is internal financial support. It is not a PPM, legal opinion, tax opinion, accounting conclusion, fairness opinion, valuation opinion, or investment recommendation. No prospective investor should receive it outside a counsel-approved offering package containing complete terms and risk disclosures.")

d.add_heading('Appendix A — Version 3 Proposed Term Register',1)
table(d,["Term","Version 3 assumption","Status"],[
 ["Maximum face principal","$3,000,000","Proposed"],["Phase 1","$1,250,000","Proposed"],["Phase 2","$1,750,000 after $1.0M annualized revenue milestone","Proposed"],["Cash coupon","4% quarterly","Proposed"],["PIK coupon","5% noncash accrual","Proposed"],["Conversion discount","20%","Illustrative"],["Valuation cap","$12 million pre-money","Illustrative"],["Qualified financing","$5 million","Illustrative"],["Maturity","November 30, 2030","Proposed"],["Forecast conversion","None through 2029","Modeling treatment"],["Section 163(j)","30% ATI stress toggle on","Planning safeguard"]
],[2.0,3.5,1.4])

d.core_properties.title="SVT Nexus 2026–2029 Financial Estimate and Offering Support Memorandum — Version 3"
d.core_properties.subject="Proposed hybrid convertible note and growth case"
d.core_properties.author="SVT Nexus"
d.core_properties.comments="Confidential draft; management, securities counsel, tax, and accounting review required"
d.save(OUT); print(OUT)
